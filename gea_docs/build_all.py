# -*- coding: utf-8 -*-
"""一键重建 GEA 全部交付物：合订本、19份制度、26份表格、简介、官网文案、总目录、ZIP。
用法：python3 build_all.py
依赖：pip install python-docx
"""
import io, os, re, sys, shutil, zipfile

BASE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, BASE)
from forms_data import FORMS

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HEI, FANG, SONG = u'黑体', u'仿宋', u'宋体'
HOME = u'/home/user'
OUT_ROOT = os.path.join(HOME, u'GEA文件包（正式版V1.0）')
DIR_ZD = os.path.join(OUT_ROOT, u'一、制度文件（19份）')
DIR_BG = os.path.join(OUT_ROOT, u'二、执行表格（26份）')
for d in (OUT_ROOT, DIR_ZD, DIR_BG):
    os.makedirs(d, exist_ok=True)

def set_font(run, name, size, bold=False):
    run.font.name = 'Times New Roman'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    run.font.size = Pt(size); run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)

def page_setup(doc):
    sec = doc.sections[0]
    sec.page_width = Cm(21); sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.54); sec.bottom_margin = Cm(2.54)
    sec.left_margin = Cm(2.8); sec.right_margin = Cm(2.6)

def add_page_number(doc):
    p = doc.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    f1 = OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'), 'begin')
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = 'PAGE'
    f2 = OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'), 'end')
    run._r.append(f1); run._r.append(it); run._r.append(f2)
    set_font(run, SONG, 9)

def make_table(doc, rows):
    tbl = doc.add_table(rows=len(rows), cols=len(rows[0]))
    tbl.style = 'Table Grid'
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            cell = tbl.cell(i, j)
            if i == 0:
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = cell.paragraphs[0].add_run(v)
            set_font(r, SONG if i else HEI, 10.5, bold=(i == 0))
    return tbl

tiao_re = re.compile(u'^(第[一二三四五六七八九十百零〇\\d]+条)　?')

content = u''
for fn in ['content_part1.md', 'content_part2.md', 'content_part3.md']:
    with io.open(os.path.join(BASE, fn), encoding='utf-8') as f:
        content += f.read() + u'\n'

# ========== 合订本 ==========
def build_combined():
    doc = Document(); page_setup(doc)
    for _ in range(5):
        doc.add_paragraph()
    for t, f, s, b in [(u'全球创业者联盟（GEA）', HEI, 26, True), (u'制度文件汇编', HEI, 36, True), (u'（正式版　V1.0）', HEI, 18, False)]:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(t); set_font(r, f, s, b)
    for _ in range(6):
        doc.add_paragraph()
    for line in [u'编制单位：全球创业者联盟总部', u'发布日期：2026年9月29日', u'密级：内部文件　仅限授权范围使用']:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(line); set_font(r, FANG, 14)
    doc.add_page_break()

    def heading(text, level):
        p = doc.add_paragraph(style='Heading %d' % level)
        p.paragraph_format.space_before = Pt(18 if level == 1 else 12)
        p.paragraph_format.space_after = Pt(10 if level == 1 else 6)
        if level == 1:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        set_font(r, HEI, {1: 18, 2: 15, 3: 13}[level], bold=(level <= 2))

    def body(text, bold_lead=None):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.first_line_indent = Cm(0.9)
        if bold_lead:
            r1 = p.add_run(bold_lead); set_font(r1, FANG, 12, True)
            r2 = p.add_run(text[len(bold_lead):]); set_font(r2, FANG, 12)
        else:
            r = p.add_run(text); set_font(r, FANG, 12)

    heading(u'编 制 说 明', 1)
    for n in [
        u'一、编纂目的。为统一全球创业者联盟（GEA）制度体系，消除多版本文件并存造成的重复、冲突与口径不一，联盟总部对既有各类制度、协议、手册及执行文件进行系统整合、去重与规范化改写，编纂形成本汇编，作为联盟各级组织和全体会员开展工作的统一依据。',
        u'二、整合范围。本汇编整合原有五十份文件，包括二十六份项目执行文件、章程制度、协议文本、分会体系、会员体系、文化荣誉体系及广州试点各版本方案中的制度性内容。',
        u'三、整合原则。去重合并、冲突消除（章程优先、新版优先、总部文件优先）、删繁就简、体例统一（编—章—节—条—款—项）。',
        u'四、文件效力层级。第一层级：联盟章程；第二层级：各项管理制度；第三层级：协议文本与26份执行文件模板。下位文件与上位文件抵触的，以上位文件为准。',
        u'五、试点方案说明。广州试点方案属阶段性执行文件，其制度性规则已纳入本汇编；时效性内容不再收录。']:
        body(n)
    doc.add_page_break()

    lines = content.split(u'\n')
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1; continue
        if line.startswith(u'### '):
            heading(line[4:].strip(), 3)
        elif line.startswith(u'## '):
            heading(line[3:].strip(), 2)
        elif line.startswith(u'# '):
            doc.add_page_break(); heading(line[2:].strip(), 1)
        elif line.startswith(u'|'):
            rows = []
            while i < len(lines) and lines[i].strip().startswith(u'|'):
                rows.append([c.strip() for c in lines[i].strip().strip(u'|').split(u'|')])
                i += 1
            make_table(doc, rows); continue
        else:
            m = tiao_re.match(line)
            body(line, m.group(1) if m else None)
        i += 1

    add_page_number(doc)
    out = os.path.join(OUT_ROOT, u'全球创业者联盟GEA制度文件汇编（合订本V1.0）.docx')
    doc.save(out)
    print('combined saved')

# ========== 独立制度 ==========
ZD_NAMES = {
 u'第一章':  (1,  u'联盟章程'), u'第二章': (2, u'组织与成员统一管理制度'),
 u'第三章':  (3,  u'总体商业运营管理制度'), u'第四章': (4, u'联盟利益与名誉保护制度'),
 u'第五章':  (5,  u'会员入会管理办法'), u'第六章': (6, u'会员考核管理办法'),
 u'第七章':  (7,  u'会员荣誉与授勋管理规范'), u'第八章': (8, u'城市分会设立与授权管理办法'),
 u'第九章':  (9,  u'分会运营管理办法'), u'第十章': (10, u'分会信用保障金管理办法'),
 u'第十一章': (11, u'项目对接闭环流程管理制度'), u'第十二章': (12, u'执行文件使用规范'),
 u'第十三章': (13, u'履约保障金管理办法'), u'第十四章': (14, u'收益结算与分润管理办法'),
 u'第十五章': (15, u'联盟合规红线与负面清单'), u'第十六章': (16, u'商业保密与数据安全管理制度'),
 u'第十七章': (17, u'争议调解管理办法'), u'第十八章': (18, u'协议体系与签署管理规范'),
 u'第十九章': (19, u'联盟文化与品牌管理规范'),
}

def doc_title(doc, title, subtitle=u'（正式版　V1.0）'):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(2)
    r = p.add_run(u'全球创业者联盟（GEA）'); set_font(r, HEI, 14)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title); set_font(r, HEI, 20, True)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(subtitle); set_font(r, HEI, 12)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(u'发布单位：全球创业者联盟总部　　发布日期：2026年9月29日')
    set_font(r, FANG, 10.5)

def build_institutions():
    chapters = []
    cur = None
    for raw in content.split(u'\n'):
        line = raw.rstrip()
        if line.startswith(u'## '):
            t = line[3:].strip()
            cur = (t.split(u'　')[0], t, [])
            chapters.append(cur)
        elif line.startswith(u'# '):
            cur = None
        elif cur is not None:
            cur[2].append(line)
    idx = []
    for key, title, body_lines in chapters:
        num, fname = ZD_NAMES[key]
        doc = Document(); page_setup(doc)
        doc_title(doc, fname)
        i = 0
        while i < len(body_lines):
            line = body_lines[i].rstrip()
            if not line.strip():
                i += 1; continue
            if line.startswith(u'### '):
                p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(6)
                r = p.add_run(line[4:].strip()); set_font(r, HEI, 13, True)
            elif line.startswith(u'|'):
                rows = []
                while i < len(body_lines) and body_lines[i].strip().startswith(u'|'):
                    rows.append([c.strip() for c in body_lines[i].strip().strip(u'|').split(u'|')])
                    i += 1
                make_table(doc, rows); continue
            else:
                m = tiao_re.match(line)
                p = doc.add_paragraph()
                p.paragraph_format.line_spacing = 1.5; p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.first_line_indent = Cm(0.9)
                if m:
                    r1 = p.add_run(m.group(1)); set_font(r1, FANG, 12, True)
                    r2 = p.add_run(line[len(m.group(1)):].lstrip(u'　')); set_font(r2, FANG, 12)
                else:
                    r = p.add_run(line); set_font(r, FANG, 12)
            i += 1
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0.9); p.paragraph_format.line_spacing = 1.5
        r = p.add_run(u'本文件由全球创业者联盟总部负责解释与修订，自发布之日起施行。'); set_font(r, FANG, 12)
        add_page_number(doc)
        doc.save(os.path.join(DIR_ZD, u'制度%02d　%s（正式版V1.0）.docx' % (num, fname)))
        idx.append((u'制度%02d' % num, fname))
    print('institutions:', len(idx))
    return idx

# ========== 独立表格 ==========
def build_forms():
    for form in FORMS:
        doc = Document(); page_setup(doc)
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(u'全球创业者联盟（GEA）'); set_font(r, HEI, 12)
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(u'文件%s　%s' % (form['num'], form['title'])); set_font(r, HEI, 16, True)
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(u'文件编号：GEA—＿＿＿＿—＿＿＿＿—%s—＿＿＿＿　　填表日期：＿＿＿＿年＿＿月＿＿日' % form['num'])
        set_font(r, SONG, 9)
        for sec in form['sections']:
            kind = sec[0]
            if kind in ('kv', 'grid', 'text'):
                p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(3)
                r = p.add_run(sec[1]); set_font(r, HEI, 11.5, True)
            if kind == 'kv':
                tbl = doc.add_table(rows=len(sec[2]), cols=2)
                tbl.style = 'Table Grid'; tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                for i, (label, val) in enumerate(sec[2]):
                    r = tbl.cell(i, 0).paragraphs[0].add_run(label); set_font(r, HEI, 10.5)
                    r = tbl.cell(i, 1).paragraphs[0].add_run(val or u''); set_font(r, SONG, 10.5)
                for row in tbl.rows:
                    row.cells[0].width = Cm(4.6); row.cells[1].width = Cm(10.8)
                    tr = row._tr; trPr = tr.get_or_add_trPr()
                    h = OxmlElement('w:trHeight'); h.set(qn('w:val'), '420'); h.set(qn('w:hRule'), 'atLeast')
                    trPr.append(h)
            elif kind == 'grid':
                header, rows = sec[2], sec[3]
                tbl = doc.add_table(rows=len(rows) + 1, cols=len(header))
                tbl.style = 'Table Grid'; tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                for j, h2 in enumerate(header):
                    c = tbl.cell(0, j); c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                    r = c.paragraphs[0].add_run(h2); set_font(r, HEI, 10, True)
                for i, row in enumerate(rows):
                    for j, v in enumerate(row):
                        r = tbl.cell(i + 1, j).paragraphs[0].add_run(v or u''); set_font(r, SONG, 10)
                    tr = tbl.rows[i + 1]._tr; trPr = tr.get_or_add_trPr()
                    h = OxmlElement('w:trHeight'); h.set(qn('w:val'), '380'); h.set(qn('w:hRule'), 'atLeast')
                    trPr.append(h)
            elif kind == 'text':
                for para in sec[2]:
                    p = doc.add_paragraph()
                    p.paragraph_format.first_line_indent = Cm(0.75)
                    p.paragraph_format.space_after = Pt(2); p.paragraph_format.line_spacing = 1.3
                    r = p.add_run(para); set_font(r, FANG, 11)
            elif kind == 'sign':
                p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(3)
                r = p.add_run(u'签署栏'); set_font(r, HEI, 11.5, True)
                parties = sec[1]
                tbl = doc.add_table(rows=2, cols=len(parties))
                tbl.style = 'Table Grid'; tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                for j, pt in enumerate(parties):
                    c = tbl.cell(0, j); c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                    r = c.paragraphs[0].add_run(pt); set_font(r, HEI, 10, True)
                    r = tbl.cell(1, j).paragraphs[0].add_run(u'\n\n签字（章）：\n\n日期：　　年　　月　　日')
                    set_font(r, SONG, 10)
            elif kind == 'note':
                for para in sec[1]:
                    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(4)
                    r = p.add_run(para); set_font(r, SONG, 9)
        add_page_number(doc)
        doc.save(os.path.join(DIR_BG, u'文件%s　%s（正式版）.docx' % (form['num'], form['title'])))
    print('forms:', len(FORMS))

# ========== 简介 ==========
def build_intro():
    doc = Document(); page_setup(doc)
    def center(text, font, size, bold=False, before=0, after=6):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after)
        r = p.add_run(text); set_font(r, font, size, bold)
    def para(text, size=12, font=FANG):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.6; p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.first_line_indent = Cm(0.9)
        r = p.add_run(text); set_font(r, font, size)
    center(u'全球创业者联盟（GEA）简介', HEI, 22, True, 6, 4)
    center(u'（正式版　V1.0）', HEI, 12, False, 0, 2)
    center(u'发布单位：全球创业者联盟总部　　发布日期：2026年9月29日', FANG, 10.5, False, 0, 14)
    center(u'让每一位创业者，都能找到一个可以信任的人', HEI, 15, True, 4, 12)
    for t in [
        u'全球创业者联盟（Global Entrepreneurs Alliance，GEA）于2019年在美国旧金山成立，是由创业者、投资人、项目方、资源方、专业服务机构与城市分会共同组成的国际商业合作组织。联盟以官方徽章为信，以「规则先于规模」为立盟之本，面向全球206个国家与地区逐步设立城市分会，共建六方共赢的跨境创业生态。',
        u'联盟以撮合为核心交付，以制度为根本保障。GEA建立了「章程—管理制度—执行文件」三级制度体系：十九项管理制度覆盖组织治理、会员管理、分会授权、合规风控与文化品牌；二十六份全球统一执行文件贯穿项目对接全流程，做到环环有文件、事事有记录、步步可追溯。',
        u'每一单撮合，都走同一套四阶段标准闭环：',
        u'——准入与登记：会员实名入会、甲乙双方资格审查，跨境需求先登记；',
        u'——评估与匹配：项目评分分级，分会先初审、总部再复核，资源智能匹配；',
        u'——履约与管理：履约保障金第三方专户托管，里程碑节点验收，风险实时预警；',
        u'——结算与复盘：服务费只在成交之后收取，收益三方签章结算，双向评价计入信用档案。',
        u'信任源自约束。联盟划定八条合规红线，触碰即除名、永不复籍；一切保障金实行第三方共管，任何机构与个人无权单方动用；争议由联盟调解机制优先化解。GEA不做只办活动的圈子，而是要建成一套能让跨境生意真正发生的制度与体系。',
        u'诚信、专业、共赢、担当——无论你在世界的哪个角落，身边总有一位可以信任的联盟伙伴。']:
        para(t)
    center(u'附：场景用语', HEI, 13, True, 14, 6)
    para(u'短版（宣传册／海报）：全球创业者联盟（GEA），2019年成立于美国旧金山，以「规则先于规模」为立盟之本，面向全球206个国家与地区设立分会。十九项制度、二十六份执行文件、保障金第三方托管、成交后才收服务费——让每一位创业者，都能找到一个可以信任的人。', 11, SONG)
    para(u'一句话版（口号／签名）：GEA——用制度让跨境生意真正发生，让每一位创业者都能找到可以信任的人。', 11, SONG)
    add_page_number(doc)
    doc.save(os.path.join(OUT_ROOT, u'全球创业者联盟GEA简介（正式版V1.0）.docx'))
    print('intro saved')

# ========== 总目录 ==========
def build_index(zd_index):
    doc = Document(); page_setup(doc)
    doc_title(doc, u'文件包总目录', u'（正式版　V1.0）')
    def sec_h(t):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(6)
        r = p.add_run(t); set_font(r, HEI, 13, True)
    sec_h(u'一、制度文件（19份）')
    rows = [[u'编号', u'制度名称']]
    for code, name in zd_index:
        rows.append([code, name + u'（正式版V1.0）'])
    make_table(doc, rows)
    sec_h(u'二、执行表格（26份）')
    rows = [[u'编号', u'表格名称']]
    for form in FORMS:
        rows.append([u'文件' + form['num'], form['title'] + u'（正式版）'])
    make_table(doc, rows)
    sec_h(u'三、其他（3份）')
    make_table(doc, [[u'名称', u'说明'],
        [u'制度文件汇编（合订本V1.0）', u'全部制度合订，供总览、培训与存档'],
        [u'GEA简介（正式版V1.0）', u'联盟官方简介，含宣传短版与口号版'],
        [u'GEA官网简介文案（正式版V1.0）', u'品牌官网五大板块文案，含编辑使用说明']])
    add_page_number(doc)
    doc.save(os.path.join(OUT_ROOT, u'00　文件包总目录（正式版V1.0）.docx'))
    print('index saved')

# ========== 打包 ==========
def build_zip():
    zip_path = os.path.join(HOME, u'GEA制度与执行文件全套（正式版V1.0）.zip')
    with zipfile.ZipFile(zip_path, 'w', zipfile.ZIP_DEFLATED) as z:
        for root, dirs, files in os.walk(OUT_ROOT):
            for f in sorted(files):
                full = os.path.join(root, f)
                z.write(full, os.path.relpath(full, os.path.dirname(OUT_ROOT)))
    print('zip:', os.path.getsize(zip_path))

if __name__ == '__main__':
    build_combined()
    zd_index = build_institutions()
    build_forms()
    build_intro()
    # 官网文案
    import make_web_intro_lib
    make_web_intro_lib.build(OUT_ROOT)
    build_index(zd_index)
    build_zip()
    print('ALL DONE')
